"""Načítání Survivor dat z .xlsx (viz schéma v README).

    from survivor_data import load_series, get_active, get_post_merge_tcs
    data = load_series('IV')   # 'III' | 'IV' | 'V'
"""
import os
import re
import openpyxl

SERIES_FILES = {
    "III": "survivor_3_data.xlsx",
    "IV": "survivor_4_data.xlsx",
    "V": "survivor_5_data.xlsx",
}
POST_MERGE_START = {"III": 1, "IV": 13, "V": 1}
# Záloha: použije se, jen když data nejdou načíst. Normálně se kola po sloučení
# auto-detekují z dat (kola s hlasy) — viz get_post_merge_tcs().
_PMT_FALLBACK = {"III": list(range(1, 9)), "IV": list(range(13, 21)), "V": [1]}


def _find_data_dir():
    """Najde složku s daty: env SURVIVOR_DATA_DIR -> ./data -> ../01_Data."""
    here = os.path.dirname(os.path.abspath(__file__))
    for d in (os.environ.get("SURVIVOR_DATA_DIR"),
              os.path.join(here, "data"),
              os.path.join(here, "..", "01_Data")):
        if d and os.path.isdir(d) and any(
                os.path.exists(os.path.join(d, f)) for f in SERIES_FILES.values()):
            return d
    return os.path.join(here, "data")


DATA_DIR = _find_data_dir()


def _split_names(cell):
    """Bunka 'Vyrazeny' -> seznam jmen. Podpora vic vyrazenych v jednom kole
    (oddelovace ; / + &). Carka se NEdeli kvuli jmenum typu 'Natalie K.'.
    Vic vyrazenych lze zadat i jako vic radku se stejnym KR."""
    if cell is None:
        return []
    return [p.strip() for p in re.split(r"[;/+&]", str(cell)) if p and p.strip()]


def load_series(series):
    wb = openpyxl.load_workbook(os.path.join(DATA_DIR, SERIES_FILES[series]), data_only=True)
    d = {"players": [], "episodes": {}, "votes": {}, "nadoby": {}, "porota": []}

    ws = wb["Hráči"]; h = [c.value for c in ws[1]]
    has_nick = "Přezdívka" in h
    nick_c = h.index("Přezdívka") if has_nick else (h.index("Jméno") if "Jméno" in h else 1)
    name_c = h.index("Plné jméno") if "Plné jméno" in h else nick_c
    pos_c = h.index("Finální pozice") if "Finální pozice" in h else 3
    jur_c = h.index("V porotě?") if "V porotě?" in h else 4
    kmen_c = h.index("Kmen před sloučením") if "Kmen před sloučením" in h else None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[name_c] is None:
            continue
        d["players"].append({"nick": r[nick_c], "name": r[name_c],
                             "final_pos": r[pos_c] if pos_c < len(r) else None,
                             "in_jury": (r[jur_c] == "ano") if jur_c < len(r) else False,
                             "kmen": (r[kmen_c] if kmen_c is not None and kmen_c < len(r) else None)})

    # Kdyz chybi sloupec 'Prezdivka' (nick = plne jmeno, napr. serie IV) -> odvod
    # kratkou prezdivku: krestni jmeno, pri kolizi 'Krestni X.' (jako JSON fallback).
    if not has_nick:
        from collections import Counter
        firsts = [str(p["nick"]).split()[0] for p in d["players"]]
        cnt = Counter(firsts)
        for p in d["players"]:
            parts = str(p["nick"]).split()
            p["nick"] = parts[0] if cnt[parts[0]] == 1 else parts[0] + " " + parts[-1][0] + "."

    n2n = {p["nick"]: p["name"] for p in d["players"]}
    n2n.update({p["name"]: p["name"] for p in d["players"]})
    norm = lambda x: n2n.get(x, x)

    if "Epizody" in wb.sheetnames:
        for r in wb["Epizody"].iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            d["episodes"][r[0]] = {"immunity": norm(r[2]) if len(r) > 2 and r[2] else None}

    ws = wb["Nádoby a duely"]; h = [c.value for c in ws[1]]; has_day = "Den" in (h or [])
    nadoba_c = h.index("Nádoba") if h and "Nádoba" in h else (3 if has_day else 2)
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        elim = r[6] if has_day else r[5]
        most = r[2] if has_day else r[1]
        nadoba = r[nadoba_c] if nadoba_c < len(r) else None
        nd = d["nadoby"].setdefault(r[0], {"most_voted": None, "eliminated": [], "nadoba": None})
        if most and not nd["most_voted"]:
            nd["most_voted"] = norm(most)
        if nadoba and not nd.get("nadoba"):
            nd["nadoba"] = str(nadoba).strip().lower()
        for nm in _split_names(elim):
            nm = norm(nm)
            if nm and nm not in nd["eliminated"]:
                nd["eliminated"].append(nm)

    for r in wb["Hlasy"].iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[1] is None or r[2] is None:
            continue
        d["votes"].setdefault(r[0], []).append((norm(r[1]), norm(r[2])))

    if "Porota" in wb.sheetnames:
        for r in wb["Porota"].iter_rows(min_row=2, values_only=True):
            if r[0] and r[1]:
                d["porota"].append((norm(r[0]), norm(r[1])))
    return d


_pmt_cache = {}


def get_post_merge_tcs(series):
    """Kola po slouceni = auto-detekce z dat (kola, kde padly hlasy).
    Pridani dalsiho kola do xlsx tedy staci - zadna editace kodu.
    Kdyz data chybi, spadne na zalozni _PMT_FALLBACK."""
    if series not in _pmt_cache:
        try:
            d = load_series(series)
            start = POST_MERGE_START[series]
            _pmt_cache[series] = sorted(t for t in d["votes"] if t >= start and d["votes"][t]) \
                or _PMT_FALLBACK.get(series, [])
        except Exception:
            _pmt_cache[series] = _PMT_FALLBACK.get(series, [])
    return _pmt_cache[series]


def get_active(data, up_to_tc, post_merge_start):
    elim = set()
    for t in range(post_merge_start, up_to_tc):
        for e in (data["nadoby"].get(t, {}).get("eliminated") or []):
            elim.add(e)
    return [p["name"] for p in data["players"] if p["name"] not in elim]
